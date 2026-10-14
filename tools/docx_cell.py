# -*- coding: utf-8 -*-
"""单元格写入辅助：在保留首段格式的前提下，安全替换整格文本。"""
import copy
from docx.oxml.ns import qn


def _clone_rpr(run):
    rpr = run._element.find(qn('w:rPr'))
    return copy.deepcopy(rpr) if rpr is not None else None


def _clone_ppr(para):
    ppr = para._p.find(qn('w:pPr'))
    return copy.deepcopy(ppr) if ppr is not None else None


def set_cell_text(cell, text):
    """把单元格正文整体替换为 text（\\n 分行），保留首段/首 run 的格式。

    正确顺序：
      1) 记录首段 pPr 与首 run rPr 作为格式模板
      2) 先删除该单元格内除首段以外的所有段落
      3) 清空首段所有 run，写入第一行
      4) 追加其余行（新段落 + 复制 pPr/rPr）
    """
    paras = cell.paragraphs
    first = paras[0]

    tmpl_rpr = _clone_rpr(first.runs[0]) if first.runs else None
    tmpl_ppr = _clone_ppr(first)

    # 2) 先删除多余段落
    for p in paras[1:]:
        p._element.getparent().remove(p._element)

    # 3) 清空首段 run
    for r in list(first.runs):
        r._element.getparent().remove(r._element)

    lines = text.split('\n') if text else ['']

    # 首行写入首段
    run = first.add_run()
    if tmpl_rpr is not None:
        run._element.insert(0, copy.deepcopy(tmpl_rpr))
    run.text = lines[0]

    # 4) 追加其余行
    for line in lines[1:]:
        new_p = copy.deepcopy(first._p)
        # 清掉复制来的 run，只保留 pPr
        for child in list(new_p):
            if child.tag == qn('w:r') or child.tag == qn('w:bookmarkStart') \
                    or child.tag == qn('w:bookmarkEnd'):
                new_p.remove(child)
        first._p.addnext(new_p)
        from docx.text.paragraph import Paragraph
        para = Paragraph(new_p, first._parent)
        r = para.add_run()
        if tmpl_rpr is not None:
            r._element.insert(0, copy.deepcopy(tmpl_rpr))
        r.text = line
        first = para
