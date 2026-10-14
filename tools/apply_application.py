# -*- coding: utf-8 -*-
"""
知穗 · 2027 挑战杯申报书 —— 内容应用脚本
================================================================
从**原始**申报书重建修订版，因此可重复运行（幂等）。

内容来源：tools/app_content.py（改文字只改那里）
输出：outputs/知穗_挑战杯申报书_修订版.docx

用法：
    python tools/apply_application.py
"""
import io
import os
import sys
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx import Document
from docx_cell import set_cell_text
import app_content as C

SRC = r'C:\Users\HUAWEI\Desktop\大挑比赛资料\知穗——面向大学生个性化AI学习助手.docx'
OUTDIR = r'C:\Users\HUAWEI\Desktop\苞米agent\outputs'
DST = os.path.join(OUTDIR, '知穗_挑战杯申报书_修订版.docx')
LOG = os.path.join(os.path.dirname(HERE), '_tmp_apply_log.txt')

AUTHORS_COVER = '冯靖豪 ，林宏亮，刘佳林'

os.makedirs(OUTDIR, exist_ok=True)
log = io.open(LOG, 'w', encoding='utf-8')
def note(m):
    log.write(m + '\n')
    print(m)


def para_set_text(p, text):
    runs = list(p.runs)
    if not runs:
        p.add_run(text)
        return
    tmpl = runs[0]
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    tmpl.text = text


shutil.copyfile(SRC, DST)
doc = Document(DST)
note('源  : %s' % SRC)
note('输出: %s' % DST)
note('')

# ---------------------------------------------------------------- 1. 封面
for p in doc.paragraphs:
    if p.text.strip().startswith('作品名称'):
        para_set_text(p, '作品名称： ' + C.PROJECT_NAME)
        note('[封面] 作品名称 -> %s' % C.PROJECT_NAME)
        break
for p in doc.paragraphs:
    if p.text.strip().startswith('申报者姓名'):
        para_set_text(p, '申报者姓名：     ' + AUTHORS_COVER)
        note('[封面] 作者顺序 -> %s' % AUTHORS_COVER)
        break

# ---------------------------------------------------------------- 2. A2 表作品名称
a2 = doc.tables[0]
set_cell_text(a2.rows[3].cells[4], C.PROJECT_NAME)
note('[A2] 作品名称已同步')

# ---------------------------------------------------------------- 3. B3 作品全称
b3 = doc.tables[1]
set_cell_text(b3.rows[0].cells[1], C.PROJECT_NAME)
note('[B3] 作品全称已同步')

# ---------------------------------------------------------------- 4. 正文各栏（来自内容文件）
for key, (ti, ri, ci) in C.TARGETS.items():
    text = C.CONTENT[key]
    cell = doc.tables[ti].rows[ri].cells[ci]
    set_cell_text(cell, text)
    note('[正文] %-11s -> 表%d r%d c%d  (%d 字)' % (key, ti, ri, ci, len(text)))

# ---------------------------------------------------------------- 5. 清除 C 节表格外遗留旧草稿
cleared = 0
for p in list(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('近年来，检索增强生成（RAG）') and '学习素养导向' in t:
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        cleared += 1
note('[清理] C 节表格外旧草稿段落：%d 处' % cleared)

# ---------------------------------------------------------------- 6. 汇总表作品名称
hit = 0
for row in doc.tables[4].rows:
    for cell in row.cells:
        if cell.text.strip().startswith('知穗'):
            set_cell_text(cell, C.PROJECT_NAME)
            hit += 1
note('[汇总表] 作品名称已同步 (%d 处)' % hit)

# ---------------------------------------------------------------- 7. D 表保持空白（用户明确指示）
note('[D 表] 按用户指示保持空白，不做填写')

doc.save(DST)
note('')
note('已保存 -> %s' % DST)
log.close()
