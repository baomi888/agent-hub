"""把 Markdown 报告转成排版好的 .docx。

用法:
    python tools/md2docx.py <input.md> [output.docx]

支持: 标题 / 正文 / 列表(含缩进) / 表格 / 代码块 / 引用 / 分隔线 / **粗体** / `代码` / 图片

图片语法(独立成行):
    ![图注文字](相对路径.png)            # 默认宽度 15.5cm
    ![图注文字](相对路径.png){w=12}      # 指定宽度 12cm
图注自动编号为 "图 N：图注文字"，居中灰字。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "微软雅黑"
CN_SERIF = "宋体"
MONO_FONT = "Consolas"

ACCENT = RGBColor(0xC8, 0x89, 0x2A)   # 玉米金，与项目主题一致
DARK = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x66, 0x66, 0x66)


def _set_font(run, name, size=None, bold=None, color=None, italic=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.font.italic = italic
    run.font.name = name
    # 中文字体必须同时设 eastasia，否则 Word 里回落到宋体
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|~~.+?~~)")


def add_rich_text(par, text, base_font, base_size, color=None):
    """解析 **粗体** / `代码` 后逐段写入 run。"""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = par.add_run(piece[2:-2])
            _set_font(r, base_font, base_size, bold=True, color=ACCENT)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            _set_font(r, MONO_FONT, base_size - 1, color=RGBColor(0xB0, 0x3A, 0x2B))
        else:
            r = par.add_run(piece)
            _set_font(r, base_font, base_size, color=color)


def add_table(doc, rows):
    header, *body = rows
    ncols = len(header)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    def fill(cells, values, bold=False, shade=None):
        for cell, val in zip(cells, values, strict=False):
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            add_rich_text(par, val.strip(), CN_FONT, 9.5,
                          color=RGBColor(0xFF, 0xFF, 0xFF) if bold else DARK)
            for r in par.runs:
                if bold:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if shade:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = tcPr.makeelement(qn("w:shd"), {})
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), shade)
                tcPr.append(shd)

    hdr = table.add_row()
    fill(hdr.cells, header, bold=True, shade="5A4632")
    for row in body:
        fill(table.add_row().cells, row)
    doc.add_paragraph()


IMG = re.compile(r"^!\[(.*?)\]\((.*?)\)(?:\{w=([\d.]+)\})?\s*$")


def add_image(doc, md_dir: Path, caption: str, rel: str, width_cm: float, fig_no: int):
    """插入居中图片 + 编号图注。"""
    img_path = (md_dir / rel).resolve()
    if not img_path.exists():
        par = doc.add_paragraph()
        r = par.add_run(f"[缺图: {rel}]")
        _set_font(r, CN_FONT, 10.5, color=RGBColor(0xB0, 0x3A, 0x2B))
        return fig_no
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.space_after = Pt(2)
    run = par.add_run()
    run.add_picture(str(img_path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    label = f"图 {fig_no}：{caption}" if caption else f"图 {fig_no}"
    add_rich_text(cap, label, CN_FONT, 9, color=GREY)
    return fig_no + 1


def convert(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    md_dir = md_path.parent

    doc = Document()
    # 页面设置：A4 + 常规页边距
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.6)
        sec.right_margin = Cm(2.6)
        sec.top_margin = Cm(2.4)
        sec.bottom_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    fig_no = 1
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # 图片（独立成行）
        m = IMG.match(line.strip())
        if m:
            caption, rel, w = m.group(1), m.group(2), m.group(3)
            fig_no = add_image(doc, md_dir, caption, rel,
                               float(w) if w else 15.5, fig_no)
            i += 1
            continue

        # 代码块
        if line.strip().startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            for b in buf:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.left_indent = Cm(0.6)
                r = par.add_run(b if b.strip() else " ")
                _set_font(r, MONO_FONT, 9, color=DARK)
            doc.add_paragraph()
            continue

        # 表格
        if line.strip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cur = lines[i].strip()
                cells = [c for c in cur.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells if c.strip()):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", line.strip()):
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(6)
            pPr = par._p.get_or_add_pPr()
            pbdr = pPr.makeelement(qn("w:pBdr"), {})
            bottom = pbdr.makeelement(qn("w:bottom"), {})
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "C8892A")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2)
            par = doc.add_paragraph()
            if level == 1:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(18)
                par.paragraph_format.space_after = Pt(10)
                add_rich_text(par, text, CN_FONT, 22, color=DARK)
                for r in par.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0x3B, 0x2C, 0x18)
            else:
                sizes = {2: 15, 3: 12.5, 4: 11, 5: 10.5, 6: 10.5}
                par.paragraph_format.space_before = Pt(12 if level == 2 else 8)
                par.paragraph_format.space_after = Pt(5)
                add_rich_text(par, text, CN_FONT, sizes.get(level, 10.5))
                for r in par.runs:
                    r.font.bold = True
                    r.font.color.rgb = ACCENT if level == 2 else RGBColor(0x5A, 0x46, 0x32)
            i += 1
            continue

        # 引用
        if line.strip().startswith(">"):
            text = re.sub(r"^>\s?", "", line.strip())
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(4)
            add_rich_text(par, text, CN_SERIF, 10, color=GREY)
            for r in par.runs:
                r.font.italic = False
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", raw)
        if m:
            indent = len(m.group(1)) // 2
            text = m.group(2)
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.left_indent = Cm(0.75 + 0.6 * indent)
            par.paragraph_format.space_after = Pt(2)
            add_rich_text(par, text, CN_FONT, 10.5)
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", raw)
        if m:
            text = m.group(3)
            par = doc.add_paragraph(style="List Number")
            par.paragraph_format.left_indent = Cm(0.75)
            par.paragraph_format.space_after = Pt(2)
            add_rich_text(par, text, CN_FONT, 10.5)
            i += 1
            continue

        # 正文
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(5)
        par.paragraph_format.line_spacing = 1.35
        par.paragraph_format.first_line_indent = Cm(0)
        add_rich_text(par, line.strip(), CN_FONT, 10.5)
        i += 1

    doc.save(str(docx_path))
    return docx_path


if __name__ == "__main__":
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    out = convert(src, dst)
    print(f"OK -> {out}")
