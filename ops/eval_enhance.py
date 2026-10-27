"""匹配器回归：拿 X3 金标集给 B2 打分（改一次词表/规则就跑一遍）。

为什么必须有这个脚本
--------------------
匹配器的错法很隐蔽：**精确率 100% 可能只是因为它几乎什么都不换**。
所以这里同时报四个数，缺一个都会误判：

  逐条全对率   94 条里有几条「换的词」和 gold 完全一致
  替换级精确率 换出去的词里有多少是对的（低了 = 乱换，会改坏检索）
  替换级召回率 gold 要求换的里面换了多少（低了 = 漏换，增强没用）
  反例误替换   把 must_not_replace 里的词给换了（**这一项必须是 0**）

用法：
    python ops/eval_enhance.py
    python ops/eval_enhance.py --vocab data/vocab/vocab_v1.csv          # X1 到货后跑这个
    python ops/eval_enhance.py --ambiguous-from-x4 docs/组员交付_2026-10-23/ambiguous_罗绍祺.xlsx
"""
from __future__ import annotations

import argparse
import csv
import os
import sys
from collections import Counter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from core.query_enhance import (  # noqa: E402
    clarify_questions,
    load_adjudication,
    load_vocab,
    match,
    norm_key,
    split_multi,
)


def load_golden(path: str) -> list[dict]:
    rows = list(csv.DictReader(open(path, encoding="utf-8-sig")))
    out = []
    for r in rows:
        q = (r.get("query") or "").strip()
        if not q:
            continue
        out.append({
            "query": q,
            "replace_to": split_multi(r.get("replace_to") or ""),
            "must_not": split_multi(r.get("must_not_replace") or ""),
        })
    return out


def ambiguous_high_risk(xlsx: str) -> set[str]:
    """从 X4 初筛表里取风险等级=高 的词，作为「需上下文裁决」集合。

    注意：X4 只是**初筛，不作终裁**。这里不是把它当真理，
    而是用它做 A/B——如果加上它召回掉一大截，说明「高歧义一律不换」这个策略
    会伤到真实需求，那就得改成「按上下文条件换」，这个结论只有数据能回答。
    """
    import openpyxl
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    header, hi = None, 0
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        if not any(cells):
            continue
        if "词" in cells and any("风险" in c for c in cells):
            header = cells
            hi = cells.index(next(c for c in cells if "风险" in c))
            break
    if header is None:
        return set()
    wi = header.index("词") if "词" in header else 1
    out = set()
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        if len(cells) <= max(hi, wi):
            continue
        if cells[hi] == "高" and cells[wi]:
            out.add(cells[wi])
    return out


def evaluate(queries: list[dict], vocab, ambiguous: set[str], verbose: int = 0,
             global_block: set[str] | None = None,
             adjudication: dict | None = None) -> dict:
    from core.query_enhance import blocked_for_query

    tp = fp = fn = 0
    exact = 0
    bad_neg: list[tuple[str, str, str]] = []   # (query, 被换的反例词, 换成了什么)
    fp_rows: list[dict] = []                   # 多换的词，带上「金标有没有声明反例」
    miss_rows: list[dict] = []
    skip_reasons: Counter = Counter()
    rewritten_text = 0                         # 检索文本把原话改掉了（必须是 0）
    rewritten_rows: list[tuple[str, str]] = []

    for item in queries:
        q = item["query"]
        # 裁决表按句生效：「有条件下入表」的词要看这一句有没有命中触发词
        blocked = set(global_block or set())
        if adjudication:
            blocked |= blocked_for_query(adjudication, q)
        r = match(q, vocab, ambiguous=ambiguous, global_block=blocked)
        got = r.canonicals
        want = item["replace_to"]

        # ★★ 「原话有没有被改」——这条指标是 2026-10-24 补上的，原因写在这里：
        # 前面四个数（全对率/精确/召回/反例）**全都测不出病句**。
        # 替换式时代它们一直是 95.7% / 97.1% / 99.0%，同时「查重率」被改成
        # 「重复率检测率」——因为指标只看"识别了哪些词"，不看"拼出来通不通"。
        # 病句是靠 X2 那 142 条真实问法人工读出来的，代价太大。
        # 现在把不变量变成指标：检索文本必须以原话为前缀，否则一律报警。
        ret = r.retrieval_text
        if not (ret == q or ret.startswith(q + " ")):
            rewritten_text += 1
            if len(rewritten_rows) < 8:
                rewritten_rows.append((q, ret))

        # 反例误替换：命中片段本身就是 must_not_replace 里的词
        neg_keys = {norm_key(n) for n in item["must_not"]}
        for h in r.hits:
            if norm_key(h.raw) in neg_keys:
                bad_neg.append((q, h.raw, h.canonical))

        got_c, want_c = Counter(got), Counter(want)
        hit_n = sum((got_c & want_c).values())
        tp += hit_n
        fp += sum(got_c.values()) - hit_n
        fn += sum(want_c.values()) - hit_n

        # ★ 把「多换的词」按金标有没有声明反例分成两类。
        # 这个区分很重要：前者是匹配器/词表的错，后者是金标漏标——
        # 两者的处理人都不一样，混在一起只会互相甩锅。
        for h in r.hits:
            if h.canonical in want_c:
                want_c[h.canonical] -= 1
                continue
            fp_rows.append({
                "query": q, "raw": h.raw, "canonical": h.canonical,
                "declared_negative": norm_key(h.raw) in neg_keys,
            })

        if Counter(got) == Counter(want):
            exact += 1
        else:
            miss_rows.append({
                "query": q, "want": want, "got": got,
                "retrieval": r.retrieval_text,
                "skipped": [(s.raw, s.reason) for s in r.skipped],
            })
        for s in r.skipped:
            skip_reasons[s.reason.split("，")[0]] += 1

    n = len(queries) or 1
    p = tp / (tp + fp) if tp + fp else 0.0
    rec = tp / (tp + fn) if tp + fn else 0.0
    undeclared = [x for x in fp_rows if not x["declared_negative"]]
    return {
        "条数": len(queries),
        "逐条全对": exact,
        "逐条全对率": exact / n,
        "替换级精确率": p,
        "替换级召回率": rec,
        "TP": tp, "FP": fp, "FN": fn,
        "反例误替换": len(bad_neg),
        "反例误替换明细": bad_neg[:8],
        "原话被改写": rewritten_text,
        "原话被改写明细": rewritten_rows,
        "多换的词": fp_rows,
        "金标未声明反例的多换": undeclared,
        "未全对的条目": miss_rows,
        "跳过原因": skip_reasons,
    }


def show(tag: str, m: dict, verbose: int) -> None:
    print(f"\n{'=' * 68}\n{tag}\n{'=' * 68}")
    print(f"  条数          {m['条数']}")
    print(f"  逐条全对      {m['逐条全对']}/{m['条数']} = {m['逐条全对率']:.1%}")
    print(f"  替换级精确率  {m['替换级精确率']:.1%}   (TP={m['TP']} FP={m['FP']})")
    print(f"  替换级召回率  {m['替换级召回率']:.1%}   (TP={m['TP']} FN={m['FN']})")
    flag = "✅" if m["反例误替换"] == 0 else "❌ 必须是 0"
    print(f"  反例误替换    {m['反例误替换']}  {flag}")
    if m["反例误替换明细"]:
        for q, raw, canon in m["反例误替换明细"]:
            print(f"      「{raw}」被换成「{canon}」  ← {q[:26]}")
    rw = m.get("原话被改写", 0)
    rflag = "✅" if rw == 0 else "❌ 增强不该动用户的原话"
    print(f"  原话被改写    {rw}  {rflag}")
    for q, ret in (m.get("原话被改写明细") or [])[:3]:
        print(f"      {q[:28]}  →  {ret[:40]}")
    und = m.get("金标未声明反例的多换") or []
    if und:
        print(f"  多换的词      {len(m['多换的词'])} 个，其中 {len(und)} 个"
              f"金标根本没声明反例（疑似漏标，不是匹配器的错）：")
        for x in und[:6]:
            print(f"      「{x['raw']}」→「{x['canonical']}」  ← {x['query'][:24]}")
    if m["跳过原因"]:
        print("  跳过原因分布：" + "、".join(f"{k} {v}" for k, v in m["跳过原因"].most_common()))
    if verbose:
        rows = m["未全对的条目"][:verbose]
        if rows:
            print(f"\n  未全对的 {len(m['未全对的条目'])} 条里，前 {len(rows)} 条：")
            for i, r in enumerate(rows, 1):
                print(f"   {i}. {r['query']}")
                print(f"      gold={r['want']}  实际={r['got']}")
                print(f"      检索文本={r['retrieval']}")
                if r["skipped"]:
                    print(f"      跳过={r['skipped']}")


def main() -> int:
    ap = argparse.ArgumentParser(description="匹配器金标回归")
    ap.add_argument("--golden", default=os.path.join(ROOT, "data/golden/golden_v1.csv"))
    ap.add_argument("--vocab", default=os.path.join(ROOT, "data/vocab/temp_from_golden.csv"))
    ap.add_argument("--ambiguous-from-x4", default="", help="X4 歧义词 xlsx（取风险=高）")
    ap.add_argument("--adjudication", default=os.path.join(ROOT, "data/vocab/adjudication.csv"),
                    help="X4 裁决表（推荐上线配置用它）")
    ap.add_argument("--verbose", type=int, default=6, help="打印多少条未全对的明细")
    args = ap.parse_args()

    vocab, warns = load_vocab(args.vocab)
    print(f"词表：{args.vocab}")
    print(f"  词条 {len(vocab)}，别名 {len(vocab.alias_to_canonical)}，反例 {len(vocab.negatives)}")
    for w in warns:
        print(f"  ⚠️ {w}")

    queries = load_golden(args.golden)
    base = evaluate(queries, vocab, set(), args.verbose)
    show("A · 不做歧义拦截（全部按词表换）", base, args.verbose)

    # ---- C 组：把金标反例当成「全局黑名单」——这是个反面教材，专门用来证明 ----
    # 为什么 X1 的 negative_aliases 必须挂在词条下，不能做成全词表禁用词。
    all_neg = {n for item in queries for n in item["must_not"] if n}
    naive = evaluate(queries, vocab, set(), 0, global_block=all_neg)
    show("C · 把反例当全局黑名单（反面教材，不要这么做）", naive, 0)
    print("\n  A→C 差值说明了什么：")
    print(f"     精确率 {base['替换级精确率']:.1%} → {naive['替换级精确率']:.1%}"
          f"　召回率 {base['替换级召回率']:.1%} → {naive['替换级召回率']:.1%}")
    print(f"     反例误替换 {base['反例误替换']} → {naive['反例误替换']}")
    print(f"     ⇒ 用 {len(all_neg)} 个「本句反例」去禁全词表，误替换确实降了，"
          f"但白丢了 {base['TP'] - naive['TP']} 处该换的")
    print("     ⇒ 结论：negative_aliases 必须**按词条挂**（term-scoped），"
          "这才是给 X1 的硬要求")

    # ---- D 组：接入裁决表（这才是推荐上线配置）----
    if args.adjudication and os.path.exists(args.adjudication):
        adj = load_adjudication(args.adjudication)
        final = evaluate(queries, vocab, set(), args.verbose, adjudication=adj)
        show("D · 接入歧义裁决表（推荐上线配置）", final, args.verbose)
        asks = [(it["query"], clarify_questions(adj, it["query"])) for it in queries]
        asks = [(q, a) for q, a in asks if a]
        print(f"\n  裁决表：{len(adj)} 词"
              f"（禁止入表 {sum(1 for a in adj.values() if a.decision == '禁止入表')}、"
              f"需澄清 {sum(1 for a in adj.values() if a.needs_clarify)}、"
              f"有条件 {sum(1 for a in adj.values() if a.is_conditional)}）")
        print(f"  这句里会触发追问的：{len(asks)}/{len(queries)} 条")
        for q, a in asks[:3]:
            print(f"      {q[:24]} → {a[0]}")
        print(f"\n  A→D：精确率 {base['替换级精确率']:.1%} → {final['替换级精确率']:.1%}"
              f"　召回率 {base['替换级召回率']:.1%} → {final['替换级召回率']:.1%}"
              f"　反例误替换 {base['反例误替换']} → {final['反例误替换']}")
    elif args.adjudication:
        print(f"\n⚠️ 找不到裁决表：{args.adjudication}（先跑 ops/build_adjudication.py）")

    if args.ambiguous_from_x4 and os.path.exists(args.ambiguous_from_x4):
        amb = ambiguous_high_risk(args.ambiguous_from_x4)
        print(f"\nX4 高风险歧义词：{len(amb)} 个")
        after = evaluate(queries, vocab, amb, 0)
        show("B · 高风险歧义词一律不换（A/B 对照）", after, 0)
        print("\n  A→B 变化（这几行决定「高歧义是否该一刀切不换」）：")
        print(f"     精确率 {base['替换级精确率']:.1%} → {after['替换级精确率']:.1%}"
              f"　召回率 {base['替换级召回率']:.1%} → {after['替换级召回率']:.1%}")
        print(f"     被歧义规则挡掉的替换 {base['TP'] - after['TP']} 处")
    elif args.ambiguous_from_x4:
        print(f"\n⚠️ 找不到 X4 文件：{args.ambiguous_from_x4}")

    print("\n" + "-" * 68)
    print("判读：精确率低 = 乱换（会改坏检索）；召回率低 = 漏换（增强白做）；")
    print("      反例误替换、原话被改写两项必须为 0，否则不能上线。")
    print("      ★ 「原话被改写」是 2026-10-24 新增的：它专门抓病句，")
    print("        前四个数对病句是瞎的（替换时代一直 97.1% 而句子已经不通）。")
    return 0 if (base["反例误替换"] == 0 and base["原话被改写"] == 0) else 1


if __name__ == "__main__":
    sys.exit(main())
